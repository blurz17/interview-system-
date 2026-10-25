import io
from pathlib import Path

import docx
import pypdf

from app.core.errors import BadRequestError


def parse_document(file_bytes: bytes, filename: str) -> str:
    ext = Path(filename).suffix.lower()

    try:
        if ext == ".pdf":
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            text_parts = []
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
            return "\n\n".join(text_parts).strip()

        elif ext == ".docx":
            doc = docx.Document(io.BytesIO(file_bytes))
            text_parts = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            text_parts.append(cell.text.strip())
            return "\n".join(text_parts).strip()

        elif ext == ".txt":
            try:
                return file_bytes.decode("utf-8").strip()
            except UnicodeDecodeError:
                return file_bytes.decode("latin-1", errors="replace").strip()

        else:
            raise BadRequestError(f"Unsupported extension for parsing: {ext}")

    except Exception as e:
        if isinstance(e, BadRequestError):
            raise
        raise BadRequestError(f"Failed to parse document: {str(e)}") from e
